"""Generate the fictional agency's sample data.

Deterministic: the same seed always writes the same rows. Writes:
  data/production/production_2026-WNN.xlsx   8 weekly production workbooks
  data/customers.csv                          40 customers
  data/quotes.csv                             quote pipeline
  evals/expected/production_2026-WNN.json     expected totals, summed here in plain
                                              Python (independent of the report script)
Then copies the canonical knowledge files into the skills that need them.

Every name, policy number and phone number is invented.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import random
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
WEEKS = list(range(36, 44))  # ISO weeks 36 to 43 of 2026
YEAR = 2026
AGENCY = "Fernhill Insurance Agency"

PRODUCERS = {
    # name: (weekly premium target, attainment range, renewal retention probability, line weights)
    "Maya Chen": (14000, (0.92, 1.18), 0.93, {"Auto": 4, "Home": 4, "Commercial": 1, "Life": 1}),
    "Luis Ortega": (12000, (0.88, 1.12), 0.92, {"Auto": 5, "Home": 2, "Commercial": 1, "Life": 1}),
    "Priya Raman": (15000, (0.95, 1.25), 0.95, {"Auto": 1, "Home": 2, "Commercial": 5, "Life": 1}),
    "Tom Becker": (10000, (0.85, 1.10), 0.90, {"Auto": 3, "Home": 3, "Commercial": 0, "Life": 3}),
    "Aisha Bello": (11000, (0.90, 1.15), 0.92, {"Auto": 3, "Home": 3, "Commercial": 1, "Life": 2}),
    "Grant Fowler": (12000, (0.52, 0.66), 0.72, {"Auto": 4, "Home": 3, "Commercial": 1, "Life": 1}),
}
PREMIUM_RANGE = {"Auto": (900, 2400), "Home": (1100, 3200), "Commercial": (2500, 9000), "Life": (600, 2500)}
PREFIX = {"Auto": "AU", "Home": "HO", "Commercial": "CM", "Life": "LF"}

FIRST = ["Maria", "James", "Elena", "Marcus", "Hannah", "Diego", "Olivia", "Samuel", "Grace", "Nathan",
         "Sofia", "Owen", "Leah", "Caleb", "Nora", "Isaac", "Ruby", "Felix", "Clara", "Jonah",
         "Ivy", "Theo", "Mila", "Adrian", "Zoe", "Victor", "Lena", "Hugo", "Alma", "Rafael"]
LAST = ["Lopez", "Whitfield", "Novak", "Okafor", "Brennan", "Castillo", "Haines", "Delgado", "Pryor", "Lindqvist",
        "Marsh", "Kowal", "Esposito", "Tanaka", "Ruiz", "Holloway", "Ferreira", "Ashby", "Mbeki", "Sorensen",
        "Quill", "Varga", "Hartley", "Moreau", "Duarte", "Albright", "Nakamura", "Petrov", "Calloway", "Ibsen"]
TOWNS = [("Fernhill", "30801"), ("Brookport", "30802"), ("Linden Bay", "30803"),
         ("Ashford Hollow", "30804"), ("Millbrook", "30805"), ("Stonebridge", "30806")]


def week_days(week: int) -> list[dt.date]:
    return [dt.date.fromisocalendar(YEAR, week, d) for d in range(1, 6)]


def make_customers(rng: random.Random) -> list[dict]:
    names, seen = [], set()
    while len(names) < 40:
        n = f"{rng.choice(FIRST)} {rng.choice(LAST)}"
        if n not in seen:
            seen.add(n)
            names.append(n)
    producers = list(PRODUCERS)
    rows = []
    for i, name in enumerate(names, start=1):
        town, zip_code = rng.choice(TOWNS)
        lines = sorted(rng.sample(["Auto", "Home", "Commercial", "Life"], k=rng.choice([1, 1, 2, 2, 3])))
        first, last = name.lower().split(" ")
        rows.append({
            "customer_id": f"C{i:03d}",
            "name": name,
            "town": town,
            "zip": zip_code,
            "phone": f"(555) 555-01{rng.randint(0, 99):02d}",
            "email": f"{first}.{last}@example.com",
            "lines_held": ", ".join(lines),
            "producer": rng.choice(producers),
            "customer_since": str(rng.randint(2012, 2026)),
        })
    return rows


def make_quotes(rng: random.Random, customers: list[dict]) -> list[dict]:
    as_of = dt.date(2026, 10, 23)
    rows = []
    for i in range(1, 31):
        c = rng.choice(customers)
        requested = as_of - dt.timedelta(days=rng.randint(0, 38))
        status = rng.choices(["Open", "Quoted", "Bound", "Lost"], weights=[4, 3, 3, 1])[0]
        last_contact = min(as_of, requested + dt.timedelta(days=rng.randint(0, 6)))
        rows.append({
            "quote_id": f"Q-{1000 + i}",
            "customer_id": c["customer_id"],
            "customer_name": c["name"],
            "line": rng.choice(["Auto", "Home", "Commercial", "Life"]),
            "producer": rng.choice(list(PRODUCERS)),
            "requested_date": requested.isoformat(),
            "status": status,
            "last_contact": last_contact.isoformat(),
        })
    return rows


def attainment_low(lo: float) -> bool:
    return lo < 0.7


def pick_line(rng: random.Random, weights: dict) -> str:
    lines = [k for k, w in weights.items() if w > 0]
    return rng.choices(lines, weights=[weights[k] for k in lines])[0]


def policy_number(rng: random.Random, line: str, used: set) -> str:
    while True:
        p = f"{PREFIX[line]}-{rng.randint(100000, 999999)}"
        if p not in used:
            used.add(p)
            return p


def make_week(rng: random.Random, week: int, customer_names: list[str], used: set) -> dict:
    days = week_days(week)
    production, renewals = [], []
    for producer, (target, (lo, hi), keep_p, weights) in PRODUCERS.items():
        goal = target * rng.uniform(lo, hi)
        written = 0.0
        # Renewals due this week for this producer.
        for _ in range(rng.randint(1, 3) if attainment_low(lo) else rng.randint(2, 5)):
            line = pick_line(rng, weights)
            lo_p, hi_p = PREMIUM_RANGE[line]
            expiring = round(rng.uniform(lo_p, hi_p), 2)
            pol = policy_number(rng, line, used)
            renewed = rng.random() < keep_p
            renewals.append({"producer": producer, "policy": pol, "line": line,
                             "expiring": expiring, "outcome": "Renewed" if renewed else "Lost"})
            if renewed:
                prem = round(expiring * rng.uniform(0.97, 1.12), 2)
                production.append({"date": rng.choice(days), "producer": producer, "policy": pol,
                                   "customer": rng.choice(customer_names), "line": line,
                                   "type": "Renewal", "premium": prem})
                written += prem
        # New business until the producer reaches this week's goal, without overshooting it much.
        misses = 0
        while written < goal and misses < 40:
            line = pick_line(rng, weights)
            lo_p, hi_p = PREMIUM_RANGE[line]
            prem = round(rng.uniform(lo_p, hi_p), 2)
            if written + prem > goal * 1.03:
                misses += 1
                continue
            production.append({"date": rng.choice(days), "producer": producer,
                               "policy": policy_number(rng, line, used),
                               "customer": rng.choice(customer_names), "line": line,
                               "type": "New", "premium": prem})
            written += prem
    production.sort(key=lambda r: (r["date"], r["producer"], r["policy"]))
    return {"production": production, "renewals": renewals}


def messy_text(rng: random.Random, value: str, p: float) -> str:
    """Return the value with the kind of noise real exports have."""
    if rng.random() >= p:
        return value
    return rng.choice([value.upper(), value.lower(), value + " ", " " + value])


def write_workbook(path: Path, week: int, data: dict, rng: random.Random, skipped_row: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Production"
    end = week_days(week)[-1]
    ws.append([f"{AGENCY} - Weekly production - week ending {end:%m/%d/%Y}"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["Date", "Producer", "Policy Number", "Customer", "Line", "Business Type", "Written Premium"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for r in data["production"]:
        premium = r["premium"]
        # Some rows come through as text with a dollar sign, as they do in real exports.
        cell_premium = f"${premium:,.2f}" if rng.random() < 0.15 else premium
        ws.append([r["date"], messy_text(rng, r["producer"], 0.1), r["policy"], r["customer"],
                   messy_text(rng, r["line"], 0.15), messy_text(rng, r["type"], 0.1), cell_premium])
        if rng.random() < 0.04:
            ws.append([])  # stray blank line
    if skipped_row:
        # A row with no premium: the report must skip it and say so.
        ws.append([week_days(week)[2], "Luis Ortega", "AU-000000", "Pending entry", "Auto", "New", None])
    ws.append([None, "TOTAL", None, None, None, None,
               round(sum(r["premium"] for r in data["production"]), 2)])
    for row in ws.iter_rows(min_row=4, max_col=1):
        for c in row:
            c.number_format = "mm/dd/yyyy"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["G"].width = 16

    rs = wb.create_sheet("Renewals Due")
    rs.append(["Producer", "Policy Number", "Line", "Expiring Premium", "Outcome"])
    for r in data["renewals"]:
        rs.append([r["producer"], r["policy"], r["line"], r["expiring"], r["outcome"]])

    ts = wb.create_sheet("Targets")
    ts.append(["Producer", "Weekly Premium Target"])
    for name, (target, *_rest) in PRODUCERS.items():
        ts.append([name, target])

    fixed = dt.datetime(2026, 10, 24, 9, 0, 0)
    wb.properties.creator = AGENCY
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.save(path)


def expected_totals(week: int, data: dict) -> dict:
    """Plain-Python totals used by the evals to check the report script."""
    prod, ren = data["production"], data["renewals"]
    by_line = {}
    for line in ["Auto", "Home", "Commercial", "Life"]:
        rows = [r for r in prod if r["line"] == line]
        by_line[line] = {"policies": len(rows), "premium": round(sum(r["premium"] for r in rows), 2)}
    renewed = sum(1 for r in ren if r["outcome"] == "Renewed")
    by_producer = {}
    for name, (target, *_rest) in PRODUCERS.items():
        prem = round(sum(r["premium"] for r in prod if r["producer"] == name), 2)
        by_producer[name] = {"premium": prem, "target": target, "pct_of_target": round(prem / target * 100, 1)}
    total = round(sum(r["premium"] for r in prod), 2)
    target_total = sum(t for t, *_ in PRODUCERS.values())
    return {
        "week": f"{YEAR}-W{week:02d}",
        "total_premium": total,
        "policies_written": len(prod),
        "new_policies": sum(1 for r in prod if r["type"] == "New"),
        "renewal_policies": sum(1 for r in prod if r["type"] == "Renewal"),
        "retention_rate": round(renewed / len(ren) * 100, 1),
        "pct_of_target": round(total / target_total * 100, 1),
        "by_line": by_line,
        "by_producer": by_producer,
    }


def sync_knowledge() -> None:
    """Skills must be self-contained, so they carry copies of the knowledge files they use."""
    kn = ROOT / "claude-project" / "knowledge"
    copies = {
        "faq.md": ["customer-reply"],
        "tone-guide.md": ["customer-reply", "weekly-owner-brief"],
        "quote-checklist.md": ["quote-intake"],
        "service-area.md": ["quote-intake"],
    }
    # The owner brief can run straight from a spreadsheet, so it carries the report scripts too.
    src = ROOT / "skills" / "production-report" / "scripts"
    for script in ("compute_report.py", "fill_narrative.py"):
        shutil.copyfile(src / script, ROOT / "skills" / "weekly-owner-brief" / "scripts" / script)
    for name, skills in copies.items():
        for skill in skills:
            dest = ROOT / "skills" / skill / "reference"
            if dest.parent.exists():
                dest.mkdir(exist_ok=True)
                shutil.copyfile(kn / name, dest / name)


def main() -> None:
    rng = random.Random(42)
    (DATA / "production").mkdir(parents=True, exist_ok=True)
    customers = make_customers(rng)
    with open(DATA / "customers.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(customers[0]))
        w.writeheader()
        w.writerows(customers)
    quotes = make_quotes(rng, customers)
    with open(DATA / "quotes.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(quotes[0]))
        w.writeheader()
        w.writerows(quotes)

    names = [c["name"] for c in customers]
    used: set = set()
    expected_dir = ROOT / "evals" / "expected"
    expected_dir.mkdir(parents=True, exist_ok=True)
    for week in WEEKS:
        data = make_week(rng, week, names, used)
        path = DATA / "production" / f"production_{YEAR}-W{week:02d}.xlsx"
        write_workbook(path, week, data, rng, skipped_row=(week == 40))
        (expected_dir / f"production_{YEAR}-W{week:02d}.json").write_text(
            json.dumps(expected_totals(week, data), indent=2) + "\n")
    sync_knowledge()
    print(f"Wrote {len(WEEKS)} production workbooks, {len(customers)} customers, {len(quotes)} quotes.")


if __name__ == "__main__":
    main()
