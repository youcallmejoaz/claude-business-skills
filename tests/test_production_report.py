"""Report totals against values worked out by hand for three small fixture weeks."""

import datetime as dt

import pytest
from openpyxl import Workbook

from agency import skill_runner

compute = skill_runner.report().compute
fill = skill_runner.narrative().fill


def workbook(path, rows, renewals, targets, title=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Production"
    if title:
        ws.append(["Sample Agency - Weekly production"])
        ws.append([])
    ws.append(["Date", "Producer", "Policy Number", "Customer", "Line", "Business Type", "Written Premium"])
    for r in rows:
        ws.append(r)
    rs = wb.create_sheet("Renewals Due")
    rs.append(["Producer", "Policy Number", "Line", "Expiring Premium", "Outcome"])
    for r in renewals:
        rs.append(r)
    ts = wb.create_sheet("Targets")
    ts.append(["Producer", "Weekly Premium Target"])
    for r in targets:
        ts.append(r)
    wb.save(path)
    return path


D = dt.date(2026, 10, 21)
TARGETS = [["Producer A", 4000], ["Producer B", 4000]]


@pytest.fixture
def week_a(tmp_path):
    # Hand totals: 1000 + 2000.50 + 500 + 4000 = 7500.50; A = 3000.50, B = 4500.
    return workbook(tmp_path / "a.xlsx", [
        [D, "Producer A", "AU-1", "X", "Auto", "New", 1000],
        [D, "Producer A", "HO-1", "X", "Home", "Renewal", "$2,000.50"],
        [],
        [D, "Producer B", "LF-1", "X", "Life", "New", 500],
        [D, "Producer B", "CM-1", "X", "Commercial", "New", 4000],
        [None, "TOTAL", None, None, None, None, 7500.50],
    ], [["Producer A", "HO-1", "Home", 1900, "Renewed"], ["Producer B", "AU-9", "Auto", 900, "Lost"],
        ["Producer A", "AU-8", "Auto", 1000, "Renewed"]], TARGETS)


@pytest.fixture
def week_b(tmp_path):
    # Messy names and a row with no premium. Hand totals: 1200 + 800 + 1500 = 3500.
    return workbook(tmp_path / "b.xlsx", [
        [D, "  producer a", "AU-2", "X", "auto ", "NEW", 1200],
        [D, "PRODUCER A", "HO-2", "X", "home", "new", "800"],
        [D, "Producer B", "AU-3", "X", "Auto", "Renewal", 1500.0],
        [D, "Producer B", "AU-4", "X", "Auto", "New", None],
    ], [["Producer B", "AU-3", "Auto", 1450, "Renewed"]], TARGETS, title=False)


@pytest.fixture
def week_c(tmp_path):
    # 9000 for A, 250 for a producer with no target. Total 9250.
    return workbook(tmp_path / "c.xlsx", [
        [D, "Producer A", "CM-2", "X", "Commercial", "New", 9000],
        [D, "Producer C", "LF-2", "X", "Life", "New", 250],
    ], [], TARGETS)


def test_week_a_totals(week_a):
    r = compute(week_a)
    t = r["totals"]
    assert t["total_premium"] == 7500.50
    assert t["policies_written"] == 4
    assert (t["new_policies"], t["renewal_policies"]) == (3, 1)
    assert t["new_premium"] == 5500 and t["renewal_premium"] == 2000.50
    assert (t["renewals_due"], t["renewals_kept"], t["renewals_lost"]) == (3, 2, 1)
    assert t["retention_rate"] == 66.7
    assert t["target_total"] == 8000
    assert t["pct_of_target"] == 93.8  # 7500.50 / 8000
    lines = {x["line"]: (x["policies"], x["premium"]) for x in r["by_line"]}
    assert lines == {"Auto": (1, 1000), "Home": (1, 2000.50), "Commercial": (1, 4000), "Life": (1, 500)}
    prod = {p["producer"]: (p["premium"], p["pct_of_target"]) for p in r["by_producer"]}
    assert prod == {"Producer A": (3000.50, 75.0), "Producer B": (4500, 112.5)}
    assert r["week"] == "2026-W43"
    assert r["source"]["rows_skipped"] == 0  # the TOTAL row and the blank row are not data issues
    assert "below_target_producer_a" in [f["id"] for f in r["flags"]]


def test_week_b_cleaning(week_b):
    r = compute(week_b)
    t = r["totals"]
    assert t["total_premium"] == 3500
    assert t["policies_written"] == 3
    assert {p["producer"]: p["policies"] for p in r["by_producer"]} == {"Producer A": 2, "Producer B": 1}
    assert {x["line"]: x["premium"] for x in r["by_line"]}["Auto"] == 2700
    assert r["source"]["rows_skipped"] == 1
    assert any("no written premium" in i for i in r["data_issues"])
    assert t["retention_rate"] == 100.0


def test_week_c_previous_week_and_unknown_producer(week_a, week_c):
    r = compute(week_c, week_a)
    t = r["totals"]
    assert t["total_premium"] == 9250
    assert t["prior_total_premium"] == 7500.50
    assert t["wow_change_pct"] == 23.3  # (9250 - 7500.50) / 7500.50
    assert t["retention_rate"] is None  # no renewals due: no made-up rate
    assert any("Producer C" in i and "no weekly target" in i for i in r["data_issues"])
    assert r["metrics"]["wow_direction"] == "up"


def test_rounding_is_half_up(week_a):
    m = compute(week_a)["metrics"]
    assert m["producer.producer_b.pct_of_target"] == "113%"  # 112.5 rounds up, not to even
    assert m["total_premium"] == "$7,501"


def test_fill_narrative_blocks_typed_numbers():
    metrics = {"total_premium": "$7,501", "pct_of_target": "94%"}
    text, problems = fill("We wrote {{total_premium}}. That is {{pct_of_target}} of target. Premium rose 12% this week.",
                          metrics)
    assert text == "We wrote $7,501. That is 94% of target."
    assert len(problems) == 1 and "12%" in problems[0]


def test_fill_narrative_blocks_unknown_placeholders():
    text, problems = fill("Revenue was {{revenue}}. Good week.", {"total_premium": "$1"})
    assert text == "Good week."
    assert "unknown placeholder" in problems[0]
